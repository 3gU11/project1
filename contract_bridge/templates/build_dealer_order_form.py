# -*- coding: utf-8 -*-
"""生成「瑞钧机械办事处订货合同」版式的订货单模板 dealer-order-form-v2.docx。

版式不再单独排版，而是直接沿用正式合同模板 dealer-contract-v1.docx（同为 A4 纵向、
同一套表线与栏位），保证系统自动生成的合同附件与纸质合同一致：

- 红框固定值写死在模板里：供方名称、单价表头「（底价）」、服务责任、其它约定事项、
  供方签署栏（单位名称、开户银行、帐号、电话、传真）；牌号固定「瑞钧」、计量单位固定「台」。
- 可变栏位换成 {{字段}} 占位符，由 contract_bridge.renderer.render_order_form 替换：
  需方、签订时间、明细（产品名称/规格型号/数量/单价）、交货时间、结算方式、
  购货单位名称（客户名）、地址、联系人、联系方式、其他特殊要求、备注。
  明细单价列默认「含税」，由映射按行填（没有机型的行留空）。
- 合同编号、签订地点、合计人民币金额（大写）、需方签署栏、随机配件留空，打印后手填。

明细固定 4 行；订单型号超过 4 种时由 renderer._expand_item_rows 复制最后一行自动加行。

用法：python build_dealer_order_form.py [--blank] [输出路径]
"""
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

BASE_TEMPLATE = Path(__file__).with_name("dealer-contract-v1.docx")
DEFAULT_OUTPUT = Path(__file__).with_name("dealer-order-form-v2.docx")

# 红框固定值（纸质合同上预先印好的内容）
FIXED_SUPPLIER = "昆山瑞钧机械科技有限公司"
FIXED_BRAND = "瑞钧"
FIXED_UNIT = "台"
FIXED_SUPPLIER_SIGNATURE = {
    "单位名称": f"单位名称：{FIXED_SUPPLIER}",
    "开户银行": "开户银行：中国工商银行昆山北门分理处",
    "帐": "帐　　号：1102 2320 0900 0027 757",
    # 供方栏宽度放不下电话与传真同一行，拆成两行避免号码被折断
    "电话": "电话：0512-57175508\n传真：0512-57175608",
}

# 段落级书签 -> 模板正文（{key} 由 _placeholder 决定是否保留为占位符）
PARAGRAPH_TEXTS = {
    "ct_supplier_header": f"供方：{FIXED_SUPPLIER}\u3000\u3000\u3000\u3000合同编号：",
    "ct_buyer_header": "需方：{dealer_name}\u3000\u3000\u3000\u3000签订地点：",
    "ct_signed_date": "一、产品名称、商标、型号、数量、金额、供货时间及数量"
                      "\u3000\u3000 签订时间：{signed_date}",
    "ct_delivery_date": "一、交（提）货时间：{delivery_date}",
    "ct_settlement_terms": "二、结算方式及期限：{settlement_terms}",
}

# 购货方详细资料表：书签 -> 字段名
BUYER_CELLS = {
    "ct_buyer_name": "buyer_name",
    "ct_buyer_address": "buyer_address",
    "ct_contact_name": "contact_name",
    "ct_contact_phone": "contact_phone",
    "ct_special_requirements": "special_requirements",
    "ct_remark": "remark",
}

ITEM_SUFFIXES = ("product", "spec", "qty")


def _placeholder(key: str, placeholders: bool) -> str:
    return "{{" + key + "}}" if placeholders else ""


def _set_paragraph_text(paragraph, text: str) -> None:
    """整体替换段落文字，保留第一个 run 的字体格式。

    空单元格没有 run 时新建的 run 必须显式用模板正文字号（10.5pt），
    否则会退回默认字号把明细行撑高、把签署栏挤到第二页。
    """
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        run = paragraph.add_run(text)
        run.font.size = Pt(10.5)


def _bookmark_names(element) -> list[str]:
    return [
        mark.get(qn("w:name"))
        for mark in element.findall(".//" + qn("w:bookmarkStart"))
        if mark.get(qn("w:name"))
    ]


def _set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if text and paragraph.runs:
        # 已有文字的单元格（台）不再重复写入
        if paragraph.text.strip() == text:
            return
    _set_paragraph_text(paragraph, text)


def build(target: Path = DEFAULT_OUTPUT, placeholders: bool = True) -> Path:
    """placeholders=True 生成渲染用模板；False 生成打印用空白合同。"""
    doc = Document(str(BASE_TEMPLATE))

    # 1) 订单头与服务条款：按书签替换整段文字
    for paragraph in doc.paragraphs:
        for name in _bookmark_names(paragraph._p):
            if name not in PARAGRAPH_TEXTS:
                continue
            text = PARAGRAPH_TEXTS[name].format(
                **{key: _placeholder(key, placeholders) for key in ("dealer_name", "signed_date",
                                                                   "delivery_date", "settlement_terms")}
            )
            _set_paragraph_text(paragraph, text)
            break

    # 2) 明细表：产品名称/规格型号/数量/单价放占位符，牌号、计量单位按纸质合同固定
    item_table = doc.tables[0]
    for index in range(1, 5):
        row_cells = item_table.rows[index].cells
        _set_cell_text(row_cells[0], _placeholder(f"item{index}_product", placeholders))
        _set_cell_text(row_cells[1], FIXED_BRAND)
        _set_cell_text(row_cells[2], _placeholder(f"item{index}_spec", placeholders))
        _set_cell_text(row_cells[3], FIXED_UNIT)
        _set_cell_text(row_cells[4], _placeholder(f"item{index}_qty", placeholders))
        _set_cell_text(row_cells[5], _placeholder(f"item{index}_price", placeholders))

    # 3) 购货方详细资料表
    buyer_table = doc.tables[1]
    for row in buyer_table.rows:
        for cell in row.cells:
            for name in _bookmark_names(cell._tc):
                key = BUYER_CELLS.get(name)
                if key:
                    _set_cell_text(cell, _placeholder(key, placeholders))
                    break

    # 4) 供方签署栏：单位名称/开户银行/帐号/电话传真写死，其余（印章、负责人）留空手填
    signature_table = doc.tables[2]
    supplier_cell = signature_table.rows[0].cells[0]
    for paragraph in supplier_cell.paragraphs:
        text = paragraph.text.strip()
        for prefix, replacement in FIXED_SUPPLIER_SIGNATURE.items():
            if text.startswith(prefix):
                _set_paragraph_text(paragraph, replacement)
                break

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    return target


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    blank = "--blank" in sys.argv
    output = Path(args[0]) if args else DEFAULT_OUTPUT
    print("saved ->", build(output, placeholders=not blank))
